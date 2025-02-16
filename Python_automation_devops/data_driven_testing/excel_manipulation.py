
#Use openpyxl for general .xlsx file operations.

# Use pandas for data analysis and manipulation.
from openpyxl import Workbook

# Create a new workbook and select the active worksheet
def test_create_and_save_workbook():
    wb = Workbook()
    ws1 = wb.active
    ws=wb.create_sheet('sample')

    # Write data to cells
    ws['A1'] = "Name"
    ws['B1'] = "Age"
    ws['A2'] = "John"
    ws['B2'] = 30

    # Save the workbook
    wb.save("example.xlsx")

def test_read_excel_file():
    from openpyxl import load_workbook
    wb=load_workbook('example.xlsx')
    ws=wb.active
    name=ws['A2'].value
    age=ws['B2'].value
    print(f'the name is {name} and age is {age}')

def test_modify_excel():
    from openpyxl import load_workbook
    wb=load_workbook('example.xlsx')
    ws=wb.active
    ws['A1']='First_name'

    wb.save('example_modify1.xlsx')

def test_iterate_excel():
    from openpyxl import load_workbook
    wb=load_workbook('example.xlsx')
    ws=wb.active
    max_row=ws.max_row
    max_col=ws.max_column
    for row in ws.iter_rows( max_col=max_col, max_row=max_row, values_only=True):
        print(row)

#....................... PANDAS ........................................
# Use pandas for data analysis and manipulation.
import pandas as pd
